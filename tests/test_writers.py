"""写入器的结构测试。"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest

from doc_converter import convert
from doc_converter.assets import AssetStore, MODE_FILE
from doc_converter.errors import ConversionError
from doc_converter.model import (CodeBlock, Document, Heading, Link, ListBlock,
                                 ListItem, Paragraph, QuoteBlock, TableBlock,
                                 TableCell, Text)
from doc_converter.pdf_engine import (PAGE_SIZES, available_engines, render_pdf)
from doc_converter.writers.html_writer import HtmlWriter, print_stylesheet
from doc_converter.writers.markdown_writer import MarkdownWriter


def _render_md(doc: Document, tmp_path: Path) -> str:
    assets = AssetStore(tmp_path, "out", mode=MODE_FILE)
    path = MarkdownWriter().write(doc, tmp_path / "out.md", assets)
    return path.read_text(encoding="utf-8")


def _render_html(doc: Document, tmp_path: Path) -> str:
    assets = AssetStore(tmp_path, "out", mode=MODE_FILE)
    return HtmlWriter().render_to_string(doc, assets)


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def test_markdown_ordered_list_start_and_nesting(tmp_path):
    doc = Document(blocks=[
        ListBlock(ordered=True, start=3, items=[
            ListItem(children=[Text("三")], blocks=[
                ListBlock(ordered=False, items=[
                    ListItem(children=[Text("子一")]),
                    ListItem(children=[Text("子二")]),
                ]),
            ]),
            ListItem(children=[Text("四")]),
        ]),
    ])
    text = _render_md(doc, tmp_path)
    assert "3. 三" in text
    assert "    - 子一" in text
    assert "4. 四" in text


def test_markdown_escapes_and_table_pipes(tmp_path):
    doc = Document(blocks=[
        Paragraph(children=[Text("1. 不是列表 *星号* [方括号] <尖括号>")]),
        TableBlock(rows=[
            [TableCell(children=[Text("A")]), TableCell(children=[Text("B|C")])],
            [TableCell(children=[Text("x|y")]),
             TableCell(children=[Text("z", code=True)])],
        ]),
    ])
    text = _render_md(doc, tmp_path)
    assert "\\*星号\\*" in text
    assert "\\[方括号\\]" in text
    assert "&lt;尖括号&gt;" not in text
    assert "\\<尖括号\\>" in text
    assert "B\\|C" in text
    assert "x\\|y" in text


def test_markdown_code_fence_with_backticks(tmp_path):
    doc = Document(blocks=[CodeBlock(text="const s = \"```\";\nnext", lang="js")])
    text = _render_md(doc, tmp_path)
    assert "````js" in text
    assert "const s" in text


def test_markdown_quote_with_blocks(tmp_path):
    doc = Document(blocks=[
        QuoteBlock(children=[Text("引用")],
                   blocks=[Paragraph(children=[Text("第二段")]),
                           ListBlock(ordered=False,
                                     items=[ListItem(children=[Text("项")])])]),
    ])
    text = _render_md(doc, tmp_path)
    assert "> 第二段" in text
    assert "> - 项" in text


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

def test_html_nested_list_colspan_and_no_banner(tmp_path):
    doc = Document(title="标题", blocks=[
        Heading(level=1, children=[Text("标题")]),
        ListBlock(ordered=True, items=[
            ListItem(children=[Text("一")], blocks=[
                ListBlock(ordered=False, items=[ListItem(children=[Text("子")])])]),
            ListItem(children=[Text("二")]),
        ]),
        TableBlock(caption="表注", rows=[
            [TableCell(children=[Text("A")], colspan=2),
             TableCell(children=[Text("B")], rowspan=2)],
            [TableCell(children=[Text("C")]), TableCell(children=[Text("D")])],
        ]),
    ])
    html = _render_html(doc, tmp_path)
    assert "<li>一<ul><li>子</li></ul></li>" in html
    assert 'colspan="2"' in html and 'rowspan="2"' in html
    assert "<caption>表注</caption>" in html
    assert "doc-meta" not in html
    assert "<title>标题</title>" in html
    assert 'name="generator"' in html


def test_html_print_stylesheet():
    css = print_stylesheet("Letter")
    assert "size: Letter" in css
    assert "@page" in css
    assert "break-inside: avoid" in css


def test_html_metadata_tags(tmp_path):
    doc = Document(title="标题", metadata={"author": "作者", "keywords": "a,b"},
                   blocks=[Paragraph(children=[Text("正文")])])
    html = _render_html(doc, tmp_path)
    assert '<meta name="author" content="作者">' in html
    assert '<meta name="keywords" content="a,b">' in html


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def test_docx_numbering_restarts_per_list(rich_md_file, tmp_path):
    """两个独立有序列表必须各自从 1 开始（不同的 numId）。"""
    source = tmp_path / "lists.md"
    source.write_text("1. A\n2. B\n\n中间段落\n\n1. C\n2. D\n", encoding="utf-8")
    result = convert(source, target="docx", output=tmp_path)
    with zipfile.ZipFile(result.output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
        numbering = archive.read("word/numbering.xml").decode("utf-8")
    num_ids = re.findall(r"<w:numId w:val=\"(\d+)\"/>", xml)
    assert len(set(num_ids)) >= 2, f"两个列表应使用不同编号：{num_ids}"
    assert numbering.count("startOverride") >= 18


def test_docx_table_merges_and_header(rich_html_file, tmp_path):
    result = convert(rich_html_file, target="docx", output=tmp_path)
    with zipfile.ZipFile(result.output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert '<w:gridSpan w:val="2"/>' in xml
    assert "<w:tblHeader/>" in xml
    from docx import Document as WordDocument

    table = WordDocument(str(result.output)).tables[0]
    assert len(table.columns) == 3


def test_docx_image_alt_and_hyperlink(rich_md_file, tmp_path):
    result = convert(rich_md_file, target="docx", output=tmp_path)
    with zipfile.ZipFile(result.output) as archive:
        document = archive.read("word/document.xml").decode("utf-8")
        rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        media = [name for name in archive.namelist()
                 if name.startswith("word/media/")]
    assert media
    assert "descr=" in document
    assert 'TargetMode="External"' in rels


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def test_pdf_engines_available():
    engines = available_engines()
    assert engines, "至少应有一个可用的 PDF 引擎"
    assert set(engines) <= {"chromium", "weasyprint", "xhtml2pdf"}


def test_pdf_render_explicit_engine(tmp_path):
    html = ("<!DOCTYPE html><html><head><meta charset='utf-8'>"
            f"<style>{print_stylesheet('A4')}</style></head>"
            "<body><h1>引擎测试</h1><p>中文内容</p></body></html>")
    engine = available_engines()[0]
    output = tmp_path / "engine.pdf"
    used = render_pdf(html, output, engine=engine, page_size="A4")
    assert used == engine
    assert output.read_bytes()[:5] == b"%PDF-"


def test_pdf_unknown_engine_raises(tmp_path):
    with pytest.raises(ConversionError):
        render_pdf("<html></html>", tmp_path / "x.pdf", engine="unknown")


def test_pdf_page_size(rich_md_file, tmp_path):
    import pdfplumber

    result = convert(rich_md_file, target="pdf", output=tmp_path,
                     page_size="Letter")
    with pdfplumber.open(str(result.output)) as pdf:
        width, height = pdf.pages[0].width, pdf.pages[0].height
    assert abs(width - 612) < 6 and abs(height - 792) < 6


def test_pdf_embeds_attachments(rich_docx_file, tmp_path):
    result = convert(rich_docx_file, target="pdf", output=tmp_path)
    from pypdf import PdfReader

    reader = PdfReader(str(result.output))
    attachments = reader.attachments
    assert any(name.endswith("sample-data.csv") for name in attachments)
    assert reader.metadata.get("/Title") or reader.metadata.get("/Creator")
