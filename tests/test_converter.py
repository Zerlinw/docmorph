# -*- coding: utf-8 -*-
"""doc-converter 测试套件。

运行::

    pytest -q
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from doc_converter import SUPPORTED_FORMATS, ConversionError, convert
from doc_converter.errors import UnsupportedFormatError

# ---------------------------------------------------------------------------
# fixtures
# ---------------------------------------------------------------------------

MARKDOWN_SOURCE = """# 测试文档

这是一个包含**加粗**、*斜体*与`行内代码`的段落，
还有一个[超链接](https://example.com)。

## 列表示例

- 第一项
- 第二项

1. 有序一
2. 有序二

## 表格

| 名称 | 值 |
| --- | --- |
| PDF | 1 |
| DOCX | 2 |

## 流程图

```mermaid
graph TD
    A["开始"] --> B["处理"]
    B["处理"] --> C["结束"]
```

## 图片

![测试图片](test.png)
"""


@pytest.fixture(scope="session")
def workdir(tmp_path_factory):
    return tmp_path_factory.mktemp("doc_converter_tests")


@pytest.fixture(scope="session")
def png_bytes() -> bytes:
    import io

    from PIL import Image

    img = Image.new("RGB", (64, 32), (30, 120, 200))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture(scope="session")
def md_file(workdir, png_bytes) -> Path:
    (workdir / "test.png").write_bytes(png_bytes)
    path = workdir / "test.md"
    path.write_text(MARKDOWN_SOURCE, encoding="utf-8")
    return path


@pytest.fixture(scope="session")
def html_file(workdir, png_bytes) -> Path:
    (workdir / "test.png").write_bytes(png_bytes)
    html = """<!DOCTYPE html>
<html lang="zh-CN">
<head><meta charset="utf-8"><title>测试文档</title></head>
<body>
<h1>测试文档</h1>
<p>这是一个包含<strong>加粗</strong>、<em>斜体</em>与<code>行内代码</code>的段落，
还有一个<a href="https://example.com">超链接</a>。</p>
<h2>列表示例</h2>
<ul><li>第一项</li><li>第二项</li></ul>
<ol><li>有序一</li><li>有序二</li></ol>
<h2>表格</h2>
<table>
<thead><tr><th>名称</th><th>值</th></tr></thead>
<tbody>
<tr><td>PDF</td><td>1</td></tr>
<tr><td>DOCX</td><td>2</td></tr>
</tbody>
</table>
<h2>图片</h2>
<p><img src="test.png" alt="测试图片"></p>
<p><a href="attachment.bin">下载附件</a></p>
</body>
</html>
"""
    path = workdir / "test.html"
    path.write_text(html, encoding="utf-8")
    (workdir / "attachment.bin").write_bytes(b"binary-attachment-content")
    return path


@pytest.fixture(scope="session")
def docx_file(workdir, png_bytes) -> Path:
    from docx import Document

    path = workdir / "test.docx"
    doc = Document()
    doc.core_properties.title = "测试文档"
    doc.add_heading("测试文档", level=1)
    doc.add_paragraph("这是一个包含加粗与斜体的段落。")
    doc.add_heading("列表示例", level=2)
    doc.add_paragraph("第一项", style="List Bullet")
    doc.add_paragraph("第二项", style="List Bullet")
    doc.add_heading("表格", level=2)
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for i, text in enumerate(("名称", "值")):
        run = table.cell(0, i).paragraphs[0].add_run(text)
        run.bold = True
    for r, row in enumerate((("PDF", "1"), ("DOCX", "2")), start=1):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    doc.add_heading("图片", level=2)
    (workdir / "test.png").write_bytes(png_bytes)
    doc.add_paragraph().add_run().add_picture(str(workdir / "test.png"))
    doc.save(str(path))
    return path


@pytest.fixture(scope="session")
def pdf_file(workdir, html_file) -> Path:
    """由 HTML 生成 PDF（自带图片与表格），用于 PDF 读取测试。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from xhtml2pdf import pisa

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    html = html_file.read_text(encoding="utf-8")
    html = html.replace(
        "</head>",
        "<style>* { font-family:'STSong-Light',sans-serif !important; }"
        "</style></head>", 1)
    out = workdir / "test.pdf"
    with open(out, "wb") as fh:
        status = pisa.CreatePDF(html.encode("utf-8"), dest=fh, encoding="utf-8")
    assert not status.err, "测试 PDF 生成失败"
    return out


# ---------------------------------------------------------------------------
# 基础能力
# ---------------------------------------------------------------------------

def test_supported_formats():
    assert set(SUPPORTED_FORMATS) == {"pdf", "docx", "html", "md"}


def test_missing_file_raises(tmp_path):
    with pytest.raises(ConversionError) as exc_info:
        convert(tmp_path / "nope.md", target="html")
    assert exc_info.value.code == "file_not_found"


def test_unsupported_format_raises(md_file):
    with pytest.raises(UnsupportedFormatError):
        convert(md_file, target="rtf")


def test_same_format_raises(md_file):
    with pytest.raises(ConversionError) as exc_info:
        convert(md_file, target="md")
    assert exc_info.value.code == "same_format"


# ---------------------------------------------------------------------------
# Markdown -> 其他
# ---------------------------------------------------------------------------

def test_md_to_html(md_file, tmp_path):
    result = convert(md_file, target="html", output=tmp_path)
    text = result.output.read_text(encoding="utf-8")
    assert "<h1>测试文档</h1>" in text
    assert "<strong>加粗</strong>" in text
    assert "<em>斜体</em>" in text
    assert "<code>行内代码</code>" in text
    assert '<a href="https://example.com">超链接</a>' in text
    assert "<table>" in text and "<th>名称</th>" in text
    assert "<li>第一项</li>" in text
    assert "mermaid" in text
    assert "test_assets" in text and ".png" in text  # 图片落盘并被引用


def test_md_to_docx(md_file, tmp_path):
    result = convert(md_file, target="docx", output=tmp_path)
    assert result.output.exists()
    with zipfile.ZipFile(result.output) as zf:
        names = zf.namelist()
        assert any(n.startswith("word/media/") for n in names), "图片应嵌入 DOCX"
        xml = zf.read("word/document.xml").decode("utf-8")
    assert "测试文档" in xml
    assert "mermaid" in xml.lower()  # 流程图附上 Mermaid 源码
    assert "<w:tbl>" in xml


def test_md_to_pdf(md_file, tmp_path):
    result = convert(md_file, target="pdf", output=tmp_path)
    assert result.output.exists()
    assert result.output.stat().st_size > 1000


# ---------------------------------------------------------------------------
# HTML -> 其他
# ---------------------------------------------------------------------------

def test_html_to_md(html_file, tmp_path):
    result = convert(html_file, target="md", output=tmp_path)
    text = result.output.read_text(encoding="utf-8")
    assert "# 测试文档" in text
    assert "**加粗**" in text
    assert "*斜体*" in text
    assert "`行内代码`" in text
    assert "[超链接](https://example.com)" in text
    assert "| 名称 | 值 |" in text
    assert "- 第一项" in text
    assert "1. 有序一" in text
    assert "![测试图片]" in text
    assert "attachment.bin" in text  # 附件链接保留


def test_html_to_docx(html_file, tmp_path):
    result = convert(html_file, target="docx", output=tmp_path)
    with zipfile.ZipFile(result.output) as zf:
        names = zf.namelist()
        assert any(n.startswith("word/media/") for n in names)
        xml = zf.read("word/document.xml").decode("utf-8")
        rels = zf.read("word/_rels/document.xml.rels").decode("utf-8")
    assert "测试文档" in xml
    assert "example.com" in rels  # 超链接关系


# ---------------------------------------------------------------------------
# DOCX -> 其他
# ---------------------------------------------------------------------------

def test_docx_to_md(docx_file, tmp_path):
    result = convert(docx_file, target="md", output=tmp_path)
    text = result.output.read_text(encoding="utf-8")
    assert "# 测试文档" in text
    assert "- 第一项" in text
    assert "| 名称 | 值 |" in text
    assert "![" in text  # 图片
    assert result.assets, "图片应保存为资源文件"
    assert all(p.exists() for p in result.assets)


def test_docx_to_html(docx_file, tmp_path):
    result = convert(docx_file, target="html", output=tmp_path)
    text = result.output.read_text(encoding="utf-8")
    assert "<h1>测试文档</h1>" in text
    assert "<li>第一项</li>" in text
    assert "<table>" in text
    assert "<img" in text


# ---------------------------------------------------------------------------
# PDF -> 其他（启发式重建）
# ---------------------------------------------------------------------------

def test_pdf_to_md(pdf_file, tmp_path):
    result = convert(pdf_file, target="md", output=tmp_path)
    text = result.output.read_text(encoding="utf-8")
    assert "测试文档" in text
    assert "|" in text  # 表格被重建
    assert result.warnings, "PDF 转换应给出启发式重建告警"


def test_pdf_to_docx(pdf_file, tmp_path):
    result = convert(pdf_file, target="docx", output=tmp_path)
    assert result.output.exists()


# ---------------------------------------------------------------------------
# 批量与 CLI
# ---------------------------------------------------------------------------

def test_batch_convert(workdir, md_file, html_file, tmp_path):
    from doc_converter import convert_batch

    results, failures = convert_batch([md_file, html_file], target="docx",
                                      output=tmp_path)
    assert not failures
    assert len(results) == 2
    assert all(r.output.exists() for r in results)


def test_batch_continues_on_error(html_file, tmp_path):
    from doc_converter import convert_batch

    results, failures = convert_batch(
        [html_file, tmp_path / "missing.md"], target="md", output=tmp_path)
    assert len(results) == 1
    assert len(failures) == 1


def test_cli_single_file(md_file, tmp_path, capsys):
    from doc_converter.cli import main

    code = main([str(md_file), "-t", "html", "-o", str(tmp_path)])
    assert code == 0
    out = capsys.readouterr().out
    assert "[OK]" in out
    assert (tmp_path / "test.html").exists()


def test_cli_missing_file(tmp_path, capsys):
    from doc_converter.cli import main

    code = main([str(tmp_path / "nope.md"), "-t", "html", "-o", str(tmp_path)])
    assert code == 1
    assert "[FAIL]" in capsys.readouterr().err


def test_cli_requires_target(md_file):
    from doc_converter.cli import main

    with pytest.raises(SystemExit):
        main([str(md_file)])


def test_cli_batch_name_conflict(md_file, html_file, tmp_path, capsys):
    """不同源格式的同名文件批量转换到同一目录时不得互相覆盖。"""
    from doc_converter.cli import main

    code = main([str(md_file), str(html_file), "-t", "docx",
                 "-o", str(tmp_path)])
    assert code == 0
    outputs = sorted(p.name for p in tmp_path.glob("*.docx"))
    assert len(outputs) == 2, f"批量输出被覆盖: {outputs}"
    assert "conflict" in capsys.readouterr().err
